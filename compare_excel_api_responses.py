import argparse
import csv
import json
import time
import urllib.request
import uuid
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

from openpyxl import load_workbook


DEFAULT_XLSX = r"C:/Users/bliuj/Desktop/kec_kparcel_productiec_tiktok_task_backup_2.xlsx"
DEFAULT_SHEET = "Result 1"
DEFAULT_URL = (
    "https://cb-tms-kp-de.kec-app.com/"
    "tms-saas-web/logistics/provider/cross_border/package_order_create"
)
DEFAULT_TOKEN = (
    "eyJ0-W1l#3RhbXAiOjE2NDg4Nzk0$zEw$DEsIm5vbmNlIjoiZHZTTUVV#TAiLCJ0b2tlbiI6"
    "Ijg0YzQzYjlkLWI3$jQtNGZk$C1hOTViLWU3NTFiNj#4ODA1YyJ9"
)
DEFAULT_REPORT = "response_compare_report.csv"
EXCEL_CELL_CHAR_LIMIT = 32767


def build_multipart_body(param_json: str) -> tuple[str, bytes]:
    boundary = "----WebKitFormBoundary" + uuid.uuid4().hex
    body = (
        f"--{boundary}\r\n"
        'Content-Disposition: form-data; name="param_json"\r\n\r\n'
        f"{param_json}\r\n"
        f"--{boundary}--\r\n"
    ).encode("utf-8")
    return boundary, body


def call_api(url: str, token: str, param_json: str, timeout: int, retries: int) -> str:
    last_error = None
    for attempt in range(retries + 1):
        try:
            boundary, body = build_multipart_body(param_json)
            request = urllib.request.Request(url, data=body, method="POST")
            request.add_header("token", token)
            request.add_header("Content-Type", f"multipart/form-data; boundary={boundary}")
            request.add_header("Content-Length", str(len(body)))
            with urllib.request.urlopen(request, timeout=timeout) as response:
                return response.read().decode("utf-8", errors="replace")
        except Exception as exc:  # noqa: BLE001
            last_error = exc
            if attempt < retries:
                time.sleep(1.5 * (attempt + 1))
    raise last_error


def minify_json_text(text: str) -> str:
    return json.dumps(json.loads(text), ensure_ascii=False, separators=(",", ":"))


def first_diff_snippet(old: str, new: str, context: int = 30) -> str:
    mismatch_at = next((i for i, (a, b) in enumerate(zip(old, new)) if a != b), None)
    if mismatch_at is None:
        mismatch_at = min(len(old), len(new))
    old_part = old[max(0, mismatch_at - context): mismatch_at + context]
    new_part = new[max(0, mismatch_at - context): mismatch_at + context]
    return (
        f"first_diff_at={mismatch_at}; "
        f"old={old_part!r}; "
        f"new={new_part!r}"
    )


def classify_response(old_response: str, new_response_text: str) -> tuple[str, int, int, str]:
    new_response_minified = minify_json_text(new_response_text)
    old_len = len(old_response)
    new_len = len(new_response_minified)
    new_prefix = new_response_minified[:old_len]

    if old_response == new_response_minified:
        return "exact_match", old_len, new_len, ""

    if old_response == new_prefix:
        if old_len == EXCEL_CELL_CHAR_LIMIT and new_len > old_len:
            return (
                "truncated_prefix_match",
                old_len,
                new_len,
                "C列内容和新响应前缀一致，但疑似被 Excel 32767 字符上限截断。",
            )
        return (
            "prefix_match_only",
            old_len,
            new_len,
            "C列内容是新响应前缀，但不是完整响应。",
        )

    return (
        "visible_mismatch",
        old_len,
        new_len,
        first_diff_snippet(old_response, new_prefix),
    )


def load_rows(xlsx_path: Path, sheet_name: str) -> list[dict]:
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    worksheet = workbook[sheet_name]
    rows = []
    for row_num, row in enumerate(
        worksheet.iter_rows(min_row=2, values_only=True),
        start=2,
    ):
        main_no, request_text, response_text = row[:3]
        if not request_text:
            continue
        rows.append(
            {
                "row_num": row_num,
                "main_no": main_no,
                "request_text": request_text,
                "response_text": response_text or "",
            }
        )
    return rows


def process_row(
    row: dict,
    url: str,
    token: str,
    timeout: int,
    retries: int,
) -> dict:
    try:
        new_response = call_api(url, token, row["request_text"], timeout, retries)
        category, old_len, new_len, note = classify_response(
            row["response_text"],
            new_response,
        )
        return {
            "row_num": row["row_num"],
            "main_no": row["main_no"],
            "category": category,
            "old_len": old_len,
            "new_len": new_len,
            "note": note,
        }
    except Exception as exc:  # noqa: BLE001
        return {
            "row_num": row["row_num"],
            "main_no": row["main_no"],
            "category": "request_error",
            "old_len": len(row["response_text"]),
            "new_len": "",
            "note": f"{type(exc).__name__}: {exc}",
        }


def write_report(results: list[dict], report_path: Path, only_diff: bool) -> None:
    diff_categories = {"visible_mismatch", "request_error"}
    rows_to_write = [
        row for row in results
        if (row["category"] in diff_categories) or not only_diff
    ]

    report_path.parent.mkdir(parents=True, exist_ok=True)
    with report_path.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(
            file,
            fieldnames=["row_num", "main_no", "category", "old_len", "new_len", "note"],
        )
        writer.writeheader()
        writer.writerows(rows_to_write)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="读取 Excel 的 B 列请求体，请求接口后与 C 列响应做比对。",
    )
    parser.add_argument("--xlsx", default=DEFAULT_XLSX, help="Excel 文件路径")
    parser.add_argument("--sheet", default=DEFAULT_SHEET, help="工作表名称")
    parser.add_argument("--url", default=DEFAULT_URL, help="接口地址")
    parser.add_argument("--token", default=DEFAULT_TOKEN, help="请求 token")
    parser.add_argument("--report", default=DEFAULT_REPORT, help="输出 CSV 路径")
    parser.add_argument("--workers", type=int, default=6, help="并发数")
    parser.add_argument("--timeout", type=int, default=90, help="单次请求超时秒数")
    parser.add_argument("--retries", type=int, default=2, help="失败重试次数")
    parser.add_argument(
        "--only-diff",
        action="store_true",
        help="只输出真正不一致或请求失败的记录",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    xlsx_path = Path(args.xlsx)
    report_path = Path(args.report)

    rows = load_rows(xlsx_path, args.sheet)
    total = len(rows)
    print(f"Loaded {total} rows from {xlsx_path}")

    results = []
    with ThreadPoolExecutor(max_workers=args.workers) as executor:
        futures = [
            executor.submit(
                process_row,
                row,
                args.url,
                args.token,
                args.timeout,
                args.retries,
            )
            for row in rows
        ]
        for index, future in enumerate(as_completed(futures), start=1):
            results.append(future.result())
            if index % 25 == 0 or index == total:
                print(f"Progress: {index}/{total}")

    results.sort(key=lambda item: item["row_num"])
    write_report(results, report_path, args.only_diff)

    summary = {}
    for item in results:
        summary[item["category"]] = summary.get(item["category"], 0) + 1

    print("Summary:")
    for key, value in sorted(summary.items()):
        print(f"  {key}: {value}")
    print(f"Report written to: {report_path.resolve()}")


if __name__ == "__main__":
    main()
