import random
import json
import time
import openpyxl
from datetime import datetime, timedelta
from locust import HttpUser, task, between


class SendTrackUser(HttpUser):
    host = "https://cb-tms-kp-de.kec-app.com"
    wait_time = between(1, 3)  # 等待时间1-3秒

    # 类变量，用于存储订单ID和当前索引
    order_ids = []
    current_index = 0
    data_loaded = False

    def on_start(self):
        """用户启动时执行，加载Excel数据"""
        if not self.data_loaded:
            self.load_order_ids()
            self.data_loaded = True

    def load_order_ids(self):
        """从Excel文件加载订单ID"""
        try:
            file_path = r'C:\Users\bliuj\Desktop\kerry\生产主流程\生产压测文件\抖音货态文件.xlsx'
            workbook = openpyxl.load_workbook(file_path)
            sheet = workbook.active

            # 假设provider_order_id在第一列，从第二行开始（第一行可能是标题）
            for row in range(2, sheet.max_row + 1):
                order_id = sheet.cell(row=row, column=1).value
                if order_id:
                    self.order_ids.append(str(order_id))

            workbook.close()
            print(f"从Excel加载了 {len(self.order_ids)} 个provider_order_id")

        except Exception as e:
            print(f"从Excel加载provider_order_id失败: {e}")
            # 如果加载失败，使用默认的订单ID
            # self.order_ids = ["ITTEST20251027002"]

    def get_next_order_id(self):
        """获取下一个订单ID，如果没有数据返回None"""
        if self.current_index < len(self.order_ids):
            order_id = self.order_ids[self.current_index]
            self.current_index += 1
            return order_id
        return None

    def generate_track_data(self, order_id):
        """根据订单ID生成轨迹数据"""
        # 生成过去24小时内的随机时间
        random_hours = random.randint(1, 24)
        operate_time = (datetime.now() - timedelta(hours=random_hours)).strftime("%Y-%m-%dT%H:%M:%S+08:00")

        track_data = {
            "provider_order_id": order_id,
            "pkg_id": order_id,  # 使用相同的订单ID
            "tracking_no": "PKXXX10400000740124000Q",
            "big_bag_no": "BOX45646115629",
            "trace_info": {
                "action_code": "cb_trucktransfer_handover",
                "operate_time": operate_time
            }
        }

        return track_data

    @task(1)
    def send_track(self):
        """发送轨迹信息"""
        # 获取下一个订单ID
        order_id = self.get_next_order_id()

        # 如果没有数据了，停止任务
        if order_id is None:
            print("所有数据已发送完毕，停止压测")
            self.stop(True)  # 停止当前用户
            return

        # 生成轨迹数据
        track_data = self.generate_track_data(order_id)

        # 准备请求头
        headers = {
            "Content-Type": "application/x-www-form-urlencoded"
        }

        # 准备表单数据
        data = {
            "param_json": json.dumps(track_data)
        }

        # 发送请求
        with self.client.post(
                "/tms-saas-web/logistics/provider/cross_border/send_track",
                data=data,
                headers=headers,
                catch_response=True,
                name="发送轨迹信息"
        ) as response:
            # 根据响应判断成功/失败
            if response.status_code == 200:
                try:
                    response_data = response.json()
                    if response_data.get("code") == 0:
                        response.success()
                        print(
                            f"轨迹发送成功: {track_data['provider_order_id']} ({self.current_index}/{len(self.order_ids)})")
                    else:
                        response.failure(f"业务失败: {response_data}")
                        print(f"轨迹发送失败: {track_data['provider_order_id']}, 错误: {response_data}")
                except:
                    response.failure("响应解析失败")
            else:
                response.failure(f"HTTP错误: {response.status_code}")